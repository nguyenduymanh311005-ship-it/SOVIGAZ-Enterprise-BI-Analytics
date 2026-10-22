import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_heading(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = 'Arial'

def add_paragraph(doc, text, bold_prefix=None, is_bullet=False):
    style = 'List Bullet' if is_bullet else 'Normal'
    p = doc.add_paragraph(style=style)
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
        p.add_run(text)
    else:
        p.add_run(text)
    for run in p.runs:
        run.font.name = 'Arial'
        run.font.size = Pt(12)

doc = Document()

style = doc.styles['Normal']
font = style.font
font.name = 'Arial'
font.size = Pt(12)

title = doc.add_heading('BÁO CÁO TỔNG KẾT CHI TIẾT DỰ ÁN SOVIGAZ', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_paragraph(doc, 'Người thực hiện: Nguyễn Duy Mạnh')

add_heading(doc, '1. Mục lục các phần trình bày', level=1)
add_paragraph(doc, 'Phần 1: Các vị trí chính trong mảng Data (Vai trò và sự phối hợp).', is_bullet=True)
add_paragraph(doc, 'Phần 2: Các bước thực hiện dự án Sovigaz (Chi tiết định nghĩa, thực thi và thách thức).', is_bullet=True)
add_paragraph(doc, 'Phần 3: Tổng kết quá trình thực tập và bài học kinh nghiệm.', is_bullet=True)

add_heading(doc, '2. Các vị trí chính trong mảng Data', level=1)
add_paragraph(doc, 'Cơ cấu nhân sự trong một dự án dữ liệu thường được phân chia thành các mắt xích chuyên biệt, phối hợp chặt chẽ để đảm bảo luồng dữ liệu thông suốt từ đầu vào đến đầu ra.')
add_paragraph(doc, ' Đi sâu vào việc trao đổi với các bên liên quan để nắm bắt nhu cầu kinh doanh. BA làm nhiệm vụ chuyển đổi từ ngôn ngữ kinh doanh sang ngôn ngữ kỹ thuật thông qua tài liệu BRD.', bold_prefix='Business Analyst (BA):', is_bullet=True)
add_paragraph(doc, ' Xây dựng luồng xử lý dữ liệu, làm sạch và chuẩn hóa dữ liệu từ hệ thống Oracle để đổ vào kho dữ liệu Data Mart nhằm tối ưu hóa truy vấn, đảm bảo tính sẵn sàng của dữ liệu.', bold_prefix='Data Engineer (DE):', is_bullet=True)
add_paragraph(doc, ' Làm việc trực tiếp với Data Mart, phân tách dữ liệu thành các bảng Fact và bảng Dimension. Sử dụng các hàm phân tích như DAX để tính toán các chỉ số phức tạp và thiết kế giao diện cho báo cáo trên Power BI.', bold_prefix='Data Analyst (DA):', is_bullet=True)
add_paragraph(doc, ' Xây dựng bộ kịch bản kiểm thử chi tiết (SIT, UAT), truy vấn dữ liệu gốc từ cơ sở dữ liệu để đối chiếu với số liệu hiển thị trên báo cáo, đảm bảo logic nghiệp vụ không bị sai lệch trước khi bàn giao.', bold_prefix='Tester (Kiểm thử viên):', is_bullet=True)

add_heading(doc, '3. Các bước thực hiện dự án Sovigaz', level=1)

add_heading(doc, '3.1. Đọc hiểu tài liệu dự án', level=2)
add_paragraph(doc, ' Nắm bắt quy mô, lĩnh vực hoạt động của doanh nghiệp Sovigaz. Phân tích ngữ cảnh kinh doanh để hiểu tại sao doanh nghiệp cần hệ thống báo cáo này và kết quả đầu ra thực sự họ mong đợi là gì.', bold_prefix='- Hiểu là gì:', is_bullet=False)
add_paragraph(doc, ' Đọc tài liệu đặc tả hệ thống, khám phá cấu trúc dữ liệu thô từ Oracle. Tìm hiểu quy trình làm việc chuẩn của toàn bộ dự án từ khâu thu thập yêu cầu đến khâu nghiệm thu.', bold_prefix='- Thực hiện:', is_bullet=False)
add_paragraph(doc, ' Gặp rào cản với các thuật ngữ chuyên ngành và luồng quy trình phức tạp. Sự bỡ ngỡ ban đầu khi tiếp cận với các công việc mang tính hệ thống cao như lập tài liệu BRD hay các chuẩn mực kiểm thử SIT, UAT thực tế.', bold_prefix='- Khó khăn:', is_bullet=False)

add_heading(doc, '3.2. Xây dựng Data Dictionary', level=2)
add_paragraph(doc, ' Là quá trình định nghĩa toàn bộ siêu dữ liệu, đóng vai trò như cuốn từ điển chuẩn hóa giúp toàn bộ nhóm dự án có chung một cách hiểu về từng trường dữ liệu, tránh sự mơ hồ.', bold_prefix='- Hiểu là gì:', is_bullet=False)
add_paragraph(doc, ' Quét qua toàn bộ các file trong Data Mart, ghi chép lại tên trường, kiểu dữ liệu, ý nghĩa nghiệp vụ, logic tính toán, và mối quan hệ giữa các trường vào bảng Excel.', bold_prefix='- Thực hiện:', is_bullet=False)
add_paragraph(doc, ' Số lượng trường dữ liệu vô cùng đồ sộ, nhiều trường có tên viết tắt hoặc định dạng phức tạp. Giải pháp vượt qua là ứng dụng mạnh mẽ AI (điển hình là công cụ Antigravity) để tự động hóa trích xuất và giải nghĩa, giúp rút ngắn hàng giờ làm việc thủ công.', bold_prefix='- Khó khăn:', is_bullet=False)

add_heading(doc, '3.3. Xây dựng BRD', level=2)
add_paragraph(doc, ' Tài liệu BRD là bản thiết kế logic của toàn bộ báo cáo, quy định rõ báo cáo sẽ có những trang nào, mỗi trang cần những biểu đồ gì, chiều phân tích và chỉ số đo lường ra sao.', bold_prefix='- Hiểu là gì:', is_bullet=False)
add_paragraph(doc, ' Dựa vào Data Dictionary, đối chiếu với nhu cầu của người dùng để hoàn thiện biểu mẫu BRD. Xác định rõ công thức toán học và logic nghiệp vụ cho từng chỉ số.', bold_prefix='- Thực hiện:', is_bullet=False)
add_paragraph(doc, ' Dữ liệu lớn đi kèm với logic tính toán các chỉ tiêu rất phức tạp. Việc dịch từ các yêu cầu kinh doanh trừu tượng sang các logic cụ thể, đo lường được là một thách thức lớn đòi hỏi tư duy phân tích sâu.', bold_prefix='- Khó khăn:', is_bullet=False)

add_heading(doc, '3.4. Xây dựng Data Model', level=2)
add_paragraph(doc, ' Tổ chức lại dữ liệu thành các bảng trung tâm chứa giao dịch và các bảng chứa thông tin mô tả theo mô hình Star Schema để Power BI có thể xử lý tối ưu nhất.', bold_prefix='- Hiểu là gì:', is_bullet=False)
add_paragraph(doc, ' Thiết lập các mối quan hệ một-nhiều, một-một, phân luồng bộ lọc chéo giữa các bảng trong Power BI.', bold_prefix='- Thực hiện:', is_bullet=False)
add_paragraph(doc, ' Việc nối bảng thường xuyên gặp các lỗi đa trị (n-n) hoặc tạo ra các mối quan hệ rất phức tạp. Bắt buộc phải rà soát kỹ lưỡng luồng chạy của dữ liệu để tránh dữ liệu bị nhân bản hoặc tính toán sai lệch.', bold_prefix='- Khó khăn:', is_bullet=False)

add_heading(doc, '3.5. Xây dựng báo cáo', level=2)
add_paragraph(doc, ' Hiện thực hóa tài liệu BRD thành một sản phẩm báo cáo trực quan, sinh động và có tính tương tác cao.', bold_prefix='- Hiểu là gì:', is_bullet=False)
add_paragraph(doc, ' Kéo thả các biểu đồ, định dạng giao diện cho chuyên nghiệp. Viết các hàm DAX để tạo các chỉ số động.', bold_prefix='- Thực hiện:', is_bullet=False)
add_paragraph(doc, ' Số lượng thẻ thông tin và biểu đồ quá nhiều, dễ gây rối mắt nếu bố cục không khoa học. Thêm vào đó, việc kết hợp các hàm DAX để xử lý logic phức tạp đôi khi làm giảm tốc độ tải của báo cáo.', bold_prefix='- Khó khăn:', is_bullet=False)

add_heading(doc, '3.6. Kiểm thử báo cáo', level=2)
add_paragraph(doc, ' Xác minh và xác nhận báo cáo hoạt động đúng như thiết kế trong BRD, và số liệu hiển thị phải khớp tuyệt đối với cơ sở dữ liệu nguồn.', bold_prefix='- Hiểu là gì:', is_bullet=False)
add_paragraph(doc, ' Viết các kịch bản kiểm thử, duyệt qua từng trường hợp cụ thể. Thực hiện việc kiểm tra chéo, chụp ảnh màn hình làm bằng chứng và ghi chú lỗi vào tệp theo dõi.', bold_prefix='- Thực hiện:', is_bullet=False)
add_paragraph(doc, ' Do số lượng chỉ tiêu cần kiểm tra cực lớn, một sai số nhỏ cũng bắt buộc người kiểm thử phải dò ngược lại từ hàm DAX, qua Data Model rồi về Data Dictionary để tìm nguyên nhân gốc rễ. Đòi hỏi tính kiên nhẫn và cẩn thận cực kỳ cao.', bold_prefix='- Khó khăn:', is_bullet=False)

add_heading(doc, '3.7. Hoàn thành và đưa ra kết luận', level=2)
add_paragraph(doc, ' Đóng gói dự án sau khi vượt qua các bài kiểm tra chất lượng.', bold_prefix='- Hiểu là gì:', is_bullet=False)
add_paragraph(doc, ' Sửa chữa các lỗi sau quá trình UAT, tối ưu hóa lại báo cáo lần cuối. Bàn giao toàn bộ tài liệu (Data Dictionary, BRD, Data Model, kịch bản kiểm thử) và tệp báo cáo hoàn chỉnh. Rút ra các bài học đúc kết.', bold_prefix='- Thực hiện:', is_bullet=False)

add_heading(doc, '4. Tổng kết quá trình thực tập', level=1)
add_paragraph(doc, ' Nắm vững toàn bộ vòng đời của một dự án Business Intelligence chuẩn mực: từ khâu phân tích nghiệp vụ, lập tài liệu, mô hình hóa đến trực quan hóa và kiểm thử. Hiểu sâu sắc sự cần thiết của việc xây dựng tài liệu lưu trữ trong việc phối hợp nhóm.', bold_prefix='- Về kiến thức chuyên môn:', is_bullet=False)
add_paragraph(doc, ' Rèn luyện thành công tư duy phân tích dữ liệu quy mô doanh nghiệp. Tăng cường khả năng xử lý, mô hình hóa những tập dữ liệu lớn, phức tạp và đôi khi thiếu tính đồng nhất.', bold_prefix='- Về xử lý dữ liệu:', is_bullet=False)
add_paragraph(doc, ' Nâng cao tư duy công nghệ qua việc làm chủ Power BI, DAX. Đặc biệt, việc tận dụng hiệu quả các trợ lý AI (như Antigravity) để xử lý dữ liệu tĩnh, tạo Data Dictionary đã giúp tối ưu hóa năng suất gấp nhiều lần so với phương pháp thủ công truyền thống.', bold_prefix='- Về công nghệ & Ứng dụng AI:', is_bullet=False)

out_dir = r"C:\Users\DELL-587K9L3\Documents\Môn học kì này\INDA\Project Sovigaz\6. Viết báo cáo và slide"
os.makedirs(out_dir, exist_ok=True)
doc.save(os.path.join(out_dir, "Bao_cao_Du_an_Sovigaz_ChiTiet.docx"))
print("Done")
